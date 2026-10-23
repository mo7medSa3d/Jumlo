#!/usr/bin/env python3
"""Independently read the positive fixtures from test-office-packages.sh.

Requires python-docx, python-pptx and Pillow. This establishes host package
readability; it does not qualify Microsoft Office or Android conversion.
"""
import argparse
from io import BytesIO
import json
from pathlib import Path

from docx import Document
from pptx import Presentation
from PIL import Image


def verify(folder: Path) -> list[dict]:
    results = []
    paragraphs = Document(folder / "paragraphs.docx")
    text = "\n".join(p.text for p in paragraphs.paragraphs)
    assert "A & B <review>\nwrapped line" in text
    assert "مرحبا بالعالم" in text
    assert "Second page" in text
    assert len(paragraphs.element.xpath(".//w:br[@w:type='page']")) == 1
    results.append({"file": "paragraphs.docx", "status": "PASS", "paragraphs": len(paragraphs.paragraphs)})

    unicode = Document(folder / "unicode.docx")
    assert unicode.paragraphs[0].text == "A�B 😀"
    results.append({"file": "unicode.docx", "status": "PASS"})

    headings = Document(folder / "headings.docx")
    assert [p.style.style_id for p in headings.paragraphs] == ["Heading1", "Normal", "Heading2"]
    assert headings.paragraphs[0].text == "عنوان المستند"
    assert headings.paragraphs[0]._p.xpath("./w:pPr/w:bidi")
    assert headings.paragraphs[0]._p.xpath("./w:r/w:rPr/w:rtl")
    results.append({"file": "headings.docx", "status": "PASS", "headings": 2, "arabic_rtl": True})

    for name, count in (("pages.pptx", 2), ("streamed.pptx", 20)):
        presentation = Presentation(folder / name)
        assert len(presentation.slides) == count
        images = 0
        for slide in presentation.slides:
            pictures = [shape for shape in slide.shapes if hasattr(shape, "image")]
            assert len(pictures) == 1
            with Image.open(BytesIO(pictures[0].image.blob)) as bitmap:
                bitmap.load()
                assert bitmap.width == 120 and bitmap.height == 180
            text_shapes = [shape for shape in slide.shapes if shape.has_text_frame]
            assert len(text_shapes) == 1 and text_shapes[0].text
            if name == "pages.pptx":
                assert "مرحبا" in text_shapes[0].text
            images += 1
        results.append({"file": name, "status": "PASS", "slides": count, "decoded_images": images})
    return results


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("folder", type=Path)
    args = parser.parse_args()
    print(json.dumps({"reader": "python-docx/python-pptx/Pillow", "results": verify(args.folder)}, ensure_ascii=False, indent=2))
